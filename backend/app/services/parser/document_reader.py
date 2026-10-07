import io
import os
import re
import zipfile
from pathlib import Path
from typing import Tuple, List, Optional
import pypdf
import docx
from app.core.exceptions import DocumentParsingError
from app.core.config import settings
from app.services.llm.prompt_defense import PromptInjectionDefense

class DocumentReader:
    MAX_PAGES = 30
    MAX_TEXT_LENGTH = 100000
    MAX_DOCX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024  # 50 MB
    MAX_ZIP_COMPRESSION_RATIO = 100

    @classmethod
    def sanitize_and_detect_hidden_text(cls, text: str) -> Tuple[str, List[str]]:
        """
        Strips hidden zero-width and invisible unicode characters, normalizes whitespace,
        defends against prompt injection attempts, and returns (cleaned_text, warnings).
        """
        warnings: List[str] = []

        # 1. Detect and strip zero-width and invisible characters
        zero_width_pattern = r"[\u200b\u200c\u200d\u200e\u200f\ufeff\u00ad]"
        if re.search(zero_width_pattern, text):
            warnings.append("Security Notice: Hidden or zero-width text detected and stripped from document.")
            cleaned = re.sub(zero_width_pattern, "", text)
        else:
            cleaned = text

        # 2. Detect prompt-injection attempts
        sanitized = PromptInjectionDefense.sanitize_untrusted_input(cleaned)
        if "[FILTERED_INSTRUCTION_ATTEMPT]" in sanitized:
            warnings.append("Security Notice: Potential prompt injection text detected and sanitized.")
        
        # 3. Normalize extra whitespace
        normalized = re.sub(r"[ \t]+", " ", sanitized)
        normalized = re.sub(r"\n{3,}", "\n\n", normalized).strip()

        # 4. Enforce max text length cap
        if len(normalized) > cls.MAX_TEXT_LENGTH:
            warnings.append(f"Notice: Document text truncated to {cls.MAX_TEXT_LENGTH} characters.")
            normalized = normalized[:cls.MAX_TEXT_LENGTH]

        return normalized, warnings

    @classmethod
    def validate_file(cls, filename: str, content: bytes) -> str:
        if not content or len(content.strip()) == 0:
            raise DocumentParsingError("Uploaded file is empty")
        
        if len(content) > settings.MAX_UPLOAD_SIZE_BYTES:
            raise DocumentParsingError(
                f"File size exceeds limit of {settings.MAX_UPLOAD_SIZE_BYTES / (1024 * 1024):.0f} MB"
            )
        
        clean_name = Path(filename).name
        ext = clean_name.rsplit(".", 1)[-1].lower() if "." in clean_name else ""
        if ext not in settings.ALLOWED_EXTENSIONS:
            raise DocumentParsingError(
                f"Unsupported file format '.{ext}'. Allowed formats: {', '.join(settings.ALLOWED_EXTENSIONS)}"
            )

        # Magic bytes validation
        if ext == "pdf":
            if not content.startswith(b"%PDF-"):
                # Check within first 1024 bytes in case of leading whitespace/BOM
                if b"%PDF-" not in content[:1024]:
                    raise DocumentParsingError("File magic bytes do not match expected PDF format.")
        elif ext == "docx":
            if not content.startswith(b"PK\x03\x04"):
                raise DocumentParsingError("File magic bytes do not match expected DOCX format (missing ZIP header).")
            # Verify valid docx zip structure and check for zip bombs
            try:
                with zipfile.ZipFile(io.BytesIO(content)) as zf:
                    filenames = zf.namelist()
                    if "word/document.xml" not in filenames:
                        raise DocumentParsingError("Invalid DOCX file: word/document.xml missing from archive.")
                    
                    total_uncompressed = sum(info.file_size for info in zf.infolist())
                    if total_uncompressed > cls.MAX_DOCX_UNCOMPRESSED_BYTES:
                        raise DocumentParsingError("Potential zip-bomb detected: DOCX uncompressed size exceeds safe limit.")
                    
                    if len(content) > 0 and (total_uncompressed / len(content)) > cls.MAX_ZIP_COMPRESSION_RATIO:
                        raise DocumentParsingError("Potential zip-bomb detected: excessive compression ratio.")

                    for info in zf.infolist():
                        if info.flag_bits & 0x1:
                            raise DocumentParsingError("Encrypted DOCX archives are not supported.")
            except zipfile.BadZipFile:
                raise DocumentParsingError("Corrupted or malformed DOCX archive.")
        elif ext == "txt":
            # Disallow executable / binary headers masquerading as text
            forbidden_headers = [b"\x7fELF", b"MZ", b"\xca\xfe\xba\xbe", b"\xfe\xed\xfa", b"PK\x03\x04"]
            for h in forbidden_headers:
                if content.startswith(h):
                    raise DocumentParsingError("Binary or executable file detected in place of plaintext.")

        return ext

    @classmethod
    def extract_text(cls, filename: str, content: bytes) -> Tuple[str, str]:
        """
        Extracts and sanitizes text from PDF, DOCX, or TXT.
        Returns: (extracted_text, file_type)
        """
        file_type = cls.validate_file(filename, content)
        
        try:
            if file_type == "pdf":
                text = cls._read_pdf(content)
            elif file_type == "docx":
                text = cls._read_docx(content)
            elif file_type == "txt":
                text = cls._read_txt(content)
            else:
                raise DocumentParsingError(f"Unsupported file type: {file_type}")
            
            cleaned, _ = cls.sanitize_and_detect_hidden_text(text)
            return cleaned, file_type
        except DocumentParsingError:
            raise
        except Exception as e:
            raise DocumentParsingError(f"Failed to extract document contents: {str(e)}")

    @classmethod
    def _read_pdf(cls, content: bytes) -> str:
        try:
            stream = io.BytesIO(content)
            reader = pypdf.PdfReader(stream)
            if reader.is_encrypted:
                try:
                    # Attempt empty password decrypt
                    decrypted = reader.decrypt("")
                    if not decrypted:
                        raise DocumentParsingError("PDF is password encrypted and cannot be processed.")
                except Exception:
                    raise DocumentParsingError("PDF is password encrypted and cannot be processed.")

            if len(reader.pages) > cls.MAX_PAGES:
                raise DocumentParsingError(f"PDF exceeds page limit ({len(reader.pages)} pages > {cls.MAX_PAGES} max).")

            pages_text = []
            for page in reader.pages:
                text = page.extract_text() or ""
                if text.strip():
                    pages_text.append(text.strip())
            
            full_text = "\n\n".join(pages_text).strip()
            if not full_text:
                raise DocumentParsingError(
                    "No readable text found in PDF. The document may be image-only (scanned), empty, or corrupted."
                )
            return full_text
        except DocumentParsingError:
            raise
        except Exception as e:
            raise DocumentParsingError(f"Malformed or unreadable PDF document: {str(e)}")

    @classmethod
    def _read_docx(cls, content: bytes) -> str:
        try:
            stream = io.BytesIO(content)
            doc = docx.Document(stream)
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            
            for table in doc.tables:
                for row in table.rows:
                    row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_cells:
                        paragraphs.append(" | ".join(row_cells))
            
            full_text = "\n".join(paragraphs).strip()
            if not full_text:
                raise DocumentParsingError("DOCX file contains no readable text.")
            return full_text
        except DocumentParsingError:
            raise
        except Exception as e:
            raise DocumentParsingError(f"Malformed DOCX file: {str(e)}")

    @classmethod
    def _read_txt(cls, content: bytes) -> str:
        for encoding in ["utf-8", "utf-8-sig", "latin-1", "cp1252"]:
            try:
                text = content.decode(encoding).strip()
                if text:
                    return text
            except UnicodeDecodeError:
                continue
        raise DocumentParsingError("Could not decode text file with standard encodings.")
